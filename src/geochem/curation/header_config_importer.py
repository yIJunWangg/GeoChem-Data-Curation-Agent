"""Import user target headers from CSV/XLSX without requiring data rows."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class HeaderConfigImporter:
    """Read the first meaningful header row from a CSV or Excel file."""

    def read_headers(self, file_path: str | Path) -> tuple[list[str], list[list[Any]]]:
        path = Path(file_path)
        suffix = path.suffix.lower()
        if suffix == ".csv":
            return self._read_csv(path)
        if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
            return self._read_xlsx(path)
        if suffix == ".xls":
            return self._read_xlsx(path)
        raise ValueError(f"Unsupported header file type: {path.suffix}")

    def _read_csv(self, path: Path) -> tuple[list[str], list[list[Any]]]:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        return self._select_header_and_samples(rows)

    def _read_xlsx(self, path: Path) -> tuple[list[str], list[list[Any]]]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            return self._select_header_and_samples(rows)
        finally:
            workbook.close()

    def _select_header_and_samples(self, rows: list[list[Any]]) -> tuple[list[str], list[list[Any]]]:
        for idx, row in enumerate(rows):
            headers = ["" if cell is None else str(cell).strip() for cell in row]
            while headers and not headers[-1]:
                headers.pop()
            non_empty = [header for header in headers if header]
            if len(non_empty) >= 1:
                samples = []
                for sample in rows[idx + 1: idx + 9]:
                    trimmed = list(sample[:len(headers)])
                    if any(cell is not None and str(cell).strip() for cell in trimmed):
                        samples.append(trimmed)
                return headers, samples
        return [], []
