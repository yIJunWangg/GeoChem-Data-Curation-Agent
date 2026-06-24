"""Script to generate test fixture files (Excel, CSV)."""

from pathlib import Path

import openpyxl
import pandas as pd


FIXTURES_DIR = Path(__file__).parent


def create_simple_excel():
    """Create a simple Excel file with major element data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Major elements"

    headers = ["Sample No.", "Latitude", "Longitude", "SiO2", "Al2O3", "Na2O", "K2O", "FeOT", "LOI"]
    ws.append(headers)

    data = [
        ["S001", 35.5, 110.2, 65.23, 15.12, 3.45, 2.87, 5.12, 1.23],
        ["S002", 35.6, 110.3, 62.18, 16.34, 3.12, 3.05, 5.67, 1.45],
        ["S003", 35.7, 110.1, 68.91, 14.23, 3.78, 2.45, 4.89, 0.98],
        ["S004", 35.4, 110.4, 59.45, 17.56, 2.89, 3.34, 6.23, 1.67],
        ["S005", 35.8, 110.0, 71.34, 13.45, 4.01, 2.12, 4.56, 0.78],
    ]
    for row in data:
        ws.append(row)

    # Unit row
    ws.append(["", "°", "°", "wt%", "wt%", "wt%", "wt%", "wt%", "wt%"])

    wb.save(FIXTURES_DIR / "table_s1_simple.xlsx")
    print("Created table_s1_simple.xlsx")


def create_multisheet_excel():
    """Create a multi-sheet Excel file."""
    wb = openpyxl.Workbook()

    # Sheet 1: Major elements
    ws1 = wb.active
    ws1.title = "Major elements"
    ws1.append(["Sample ID", "SiO2", "Al2O3", "Na2O", "K2O", "CaO", "MgO"])
    ws1.append(["", "wt%", "wt%", "wt%", "wt%", "wt%", "wt%"])
    ws1.append(["T001", 63.21, 15.67, 3.23, 2.56, 4.12, 2.34])
    ws1.append(["T002", 60.45, 16.89, 2.98, 2.78, 4.56, 2.67])

    # Sheet 2: Trace elements
    ws2 = wb.create_sheet("Trace elements")
    ws2.append(["Sample ID", "Sr", "Ba", "Rb", "Zr", "Nb", "Y"])
    ws2.append(["", "ppm", "ppm", "ppm", "ppm", "ppm", "ppm"])
    ws2.append(["T001", 345, 567, 89, 234, 12, 28])
    ws2.append(["T002", 312, 623, 95, 256, 14, 31])

    # Sheet 3: Metadata
    ws3 = wb.create_sheet("Sample info")
    ws3.append(["Sample ID", "Latitude", "Longitude", "Formation", "Lithology"])
    ws3.append(["T001", 36.12, 112.34, "Shanxi Fm", "Sandstone"])
    ws3.append(["T002", 36.15, 112.38, "Shanxi Fm", "Siltstone"])

    wb.save(FIXTURES_DIR / "table_s2_multisheet.xlsx")
    print("Created table_s2_multisheet.xlsx")


def create_csv():
    """Create a CSV file with geochemical data."""
    data = {
        "Sample": ["R001", "R002", "R003", "R004"],
        "SiO₂": [64.5, 61.2, 67.8, 58.9],
        "Al₂O₃": [14.8, 16.1, 13.9, 17.2],
        "Na (wt%)": [2.34, 2.12, 2.56, 1.98],
        "K (wt%)": [2.89, 3.12, 2.45, 3.34],
        "CIA": [58.3, 62.1, 55.7, 64.5],
    }
    df = pd.DataFrame(data)
    df.to_csv(FIXTURES_DIR / "geochem_sample.csv", index=False)
    print("Created geochem_sample.csv")


if __name__ == "__main__":
    create_simple_excel()
    create_multisheet_excel()
    create_csv()
    print("All fixtures created.")
