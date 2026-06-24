"""Tests for ExcelReader."""

from pathlib import Path

import pytest

from geochem.extractors.excel_reader import ExcelReader

FIXTURES = Path(__file__).parent.parent / "fixtures"


@pytest.fixture
def reader():
    return ExcelReader()


def test_read_xinmen2022(reader):
    """Read the real ordovician Excel fixture."""
    results = reader.read(FIXTURES / "xinmen2022_ordovician.xlsx")

    assert len(results) == 1
    content = results[0]

    assert content.sheet_name == "XinMen2022"
    assert content.row_count == 21
    assert len(content.headers) > 100  # ~156 columns
    assert content.source_type == "excel"
    assert "XD2P-B26" in content.text  # First sample ID


def test_header_unit_extraction(reader):
    """Test unit extraction from various header formats."""
    # Parenthesized unit
    clean, unit = reader._extract_header_unit("Al2O3(wt%)")
    assert clean == "Al2O3"
    assert unit == "wt%"

    # Newline-separated unit
    clean, unit = reader._extract_header_unit("Li\nppm")
    assert clean == "Li"
    assert unit == "ppm"

    # Isotope with unit
    clean, unit = reader._extract_header_unit("δ15Nbulk\n‰")
    assert clean == "δ15Nbulk"
    assert unit == "‰"

    # Complex unit
    clean, unit = reader._extract_header_unit("OI\nmg CO2/g TOC")
    assert clean == "OI"
    assert unit == "mg CO2/g TOC"

    # Slash-separated
    clean, unit = reader._extract_header_unit("Depth/m")
    assert clean == "Depth"
    assert unit == "m"

    # No unit (ratio)
    clean, unit = reader._extract_header_unit("Fepy/Fehr")
    assert clean == "Fepy/Fehr"
    assert unit == ""

    # No unit (index)
    clean, unit = reader._extract_header_unit("CIA")
    assert clean == "CIA"
    assert unit == ""

    # Empty header
    clean, unit = reader._extract_header_unit("")
    assert clean == ""
    assert unit == ""


def test_header_units_dict(reader):
    """Verify header_units dict is populated correctly."""
    results = reader.read(FIXTURES / "xinmen2022_ordovician.xlsx")
    content = results[0]

    # Should have extracted some units
    assert len(content.header_units) > 0
    # Check specific units exist
    assert "wt%" in content.header_units.values() or "‰" in content.header_units.values()


def test_markdown_format(reader):
    """Verify markdown table format."""
    results = reader.read(FIXTURES / "xinmen2022_ordovician.xlsx")
    text = results[0].text

    # Should have markdown table structure
    assert "| " in text
    assert " |" in text
    assert "---" in text
    # Should have header + separator + data rows
    lines = text.split("\n")
    assert len(lines) >= 23  # 1 header + 1 separator + 21 data rows


def test_raw_headers_preserved(reader):
    """Verify raw_headers are preserved (with units)."""
    results = reader.read(FIXTURES / "xinmen2022_ordovician.xlsx")
    content = results[0]

    # raw_headers should differ from headers (contain units)
    assert len(content.raw_headers) == len(content.headers)
    # At least some raw_headers should contain \n or ()
    has_unit = any("\n" in h or "(" in h for h in content.raw_headers if h)
    assert has_unit


def test_raw_rows_populated(reader):
    """Verify raw_rows is populated for programmatic extraction."""
    results = reader.read(FIXTURES / "xinmen2022_ordovician.xlsx")
    content = results[0]

    assert len(content.raw_rows) == 21  # 21 data rows
    # Each row should have same length as headers
    for row in content.raw_rows:
        assert len(row) == len(content.headers)
    # First row should have a SampleID
    assert content.raw_rows[0][0] is not None


def test_empty_file(tmp_path):
    """Empty Excel file returns no content."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Empty"
    path = tmp_path / "empty.xlsx"
    wb.save(path)
    wb.close()

    reader = ExcelReader()
    results = reader.read(path)
    assert results == []


def test_multi_sheet(tmp_path):
    """Multi-sheet Excel returns multiple RawContent."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["A", "B"])
    ws1.append([1, 2])

    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["X", "Y"])
    ws2.append([3, 4])

    path = tmp_path / "multi.xlsx"
    wb.save(path)
    wb.close()

    reader = ExcelReader()
    results = reader.read(path)
    assert len(results) == 2
    assert results[0].sheet_name == "Sheet1"
    assert results[1].sheet_name == "Sheet2"
